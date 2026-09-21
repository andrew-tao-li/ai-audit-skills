#!/usr/bin/env python3
"""Offline procurement red-flag scanner with traceable evidence."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import statistics
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

VERSION = "0.1.1"
SKILL = "procurement-fraud"

SCHEMAS = {
    "vendors": {
        "required": ["vendor_id", "vendor_name"],
        "fields": {
            "vendor_id": ["vendor_id", "supplier_id", "供应商编码", "供应商编号"],
            "vendor_name": ["vendor_name", "supplier_name", "供应商名称", "供应商"],
            "legal_name": ["legal_name", "company_name", "企业名称", "法定名称"],
            "tax_id": ["tax_id", "tax_number", "统一社会信用代码", "税号"],
            "bank_account": ["bank_account", "bank_no", "银行账号", "收款账号"],
            "phone": ["phone", "telephone", "联系电话", "手机号"],
            "email": ["email", "邮箱"],
            "address": ["address", "registered_address", "地址", "注册地址"],
            "legal_representative": ["legal_representative", "owner_name", "法人", "法定代表人"],
            "created_at": ["created_at", "create_date", "创建日期"],
        },
    },
    "purchase_orders": {
        "required": ["po_id", "vendor_id", "buyer_id", "unit_price", "total_amount", "order_date"],
        "fields": {
            "po_id": ["po_id", "order_id", "采购订单号", "订单号"],
            "vendor_id": ["vendor_id", "supplier_id", "供应商编码", "供应商编号"],
            "buyer_id": ["buyer_id", "purchaser_id", "采购员", "采购人员编号"],
            "category": ["category", "采购类别", "品类"],
            "item": ["item", "material", "物料", "商品", "项目"],
            "unit": ["unit", "计量单位", "单位"],
            "region": ["region", "地区", "区域"],
            "quantity": ["quantity", "qty", "数量"],
            "unit_price": ["unit_price", "price", "单价"],
            "total_amount": ["total_amount", "amount", "订单金额", "总金额"],
            "currency": ["currency", "币种"],
            "order_date": ["order_date", "po_date", "订单日期", "采购日期"],
            "approval_date": ["approval_date", "approved_at", "审批日期"],
            "receipt_date": ["receipt_date", "received_at", "收货日期", "验收日期"],
        },
    },
    "employees": {
        "required": ["employee_id"],
        "fields": {
            "employee_id": ["employee_id", "emp_id", "员工编号", "工号"],
            "employee_name": ["employee_name", "name", "员工姓名", "姓名"],
            "department": ["department", "dept", "部门"],
            "phone": ["phone", "联系电话", "手机号"],
            "email": ["email", "邮箱"],
            "address": ["address", "地址", "家庭地址"],
            "bank_account": ["bank_account", "银行账号", "工资卡号"],
        },
    },
    "payments": {
        "required": ["payment_id", "po_id", "vendor_id", "amount", "payment_date"],
        "fields": {
            "payment_id": ["payment_id", "付款编号", "支付编号"],
            "po_id": ["po_id", "order_id", "采购订单号", "订单号"],
            "vendor_id": ["vendor_id", "supplier_id", "供应商编码"],
            "amount": ["amount", "payment_amount", "付款金额"],
            "currency": ["currency", "币种"],
            "payment_date": ["payment_date", "paid_at", "付款日期", "支付日期"],
            "bank_account": ["bank_account", "银行账号", "收款账号"],
        },
    },
    "bids": {
        "required": ["tender_id", "lot_id", "bidder_id", "bid_price"],
        "fields": {
            "tender_id": ["tender_id", "招标编号", "项目编号"],
            "lot_id": ["lot_id", "标段编号", "包件编号"],
            "bidder_id": ["bidder_id", "vendor_id", "投标人编号", "供应商编码"],
            "bid_price": ["bid_price", "price", "投标报价", "报价"],
            "currency": ["currency", "币种"],
            "document_path": ["document_path", "file_path", "投标文件路径", "文本路径"],
            "submitted_at": ["submitted_at", "submission_time", "提交时间"],
        },
    },
}
FILE_CANDIDATES = {
    "vendors": ["vendors.csv", "vendors.xlsx", "vendor_master.csv", "vendor_master.xlsx"],
    "purchase_orders": ["purchase_orders.csv", "purchase_orders.xlsx", "pos.csv", "pos.xlsx"],
    "employees": ["employees.csv", "employees.xlsx"],
    "payments": ["payments.csv", "payments.xlsx"],
    "bids": ["bids.csv", "bids.xlsx", "tenders.csv", "tenders.xlsx"],
}
NUMBER_FIELDS = {"quantity", "unit_price", "total_amount", "amount", "bid_price"}
DATE_FIELDS = {"created_at", "order_date", "approval_date", "receipt_date", "payment_date"}


def text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip()


def header_key(value: Any) -> str:
    return re.sub(r"[\s_\-./\\]+", "", text(value).lower())


def match_key(value: Any) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff@]+", "", text(value).lower())


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def number(value: Any) -> Optional[float]:
    raw = text(value)
    if not raw:
        return None
    negative = raw.startswith("(") and raw.endswith(")")
    cleaned = re.sub(r"[^0-9.+-]", "", raw.replace(",", ""))
    try:
        result = float(cleaned)
        if not math.isfinite(result):
            return None
        return -result if negative else result
    except ValueError:
        return None


def date_value(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = text(value)
    if not raw:
        return None
    raw = raw.split("T")[0].split(" ")[0]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def load_json(path: Optional[Path]) -> Dict[str, Any]:
    if path is None:
        return {}
    with path.open("r", encoding="utf-8-sig") as handle:
        result = json.load(handle)
    if not isinstance(result, dict):
        raise ValueError("config 顶层必须是 JSON 对象")
    return result


def read_table(path: Path) -> Tuple[List[str], List[Dict[str, Any]], Optional[str]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = [text(h) for h in (reader.fieldnames or [])]
            rows = []
            for row_number, source in enumerate(reader, start=2):
                row = {text(k): v for k, v in source.items() if k is not None}
                row.update({"_source_file": path.name, "_source_row": row_number, "_source_sheet": None})
                rows.append(row)
            return headers, rows, None
    if path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("读取 XLSX 需要 openpyxl；请离线安装或另存为 UTF-8 CSV") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        values = worksheet.iter_rows(values_only=True)
        first = next(values, None)
        if first is None:
            return [], [], worksheet.title
        headers = [text(value) for value in first]
        rows = []
        for row_number, values_row in enumerate(values, start=2):
            row = {headers[i]: value for i, value in enumerate(values_row) if i < len(headers)}
            row.update({"_source_file": path.name, "_source_row": row_number, "_source_sheet": worksheet.title})
            rows.append(row)
        workbook.close()
        return headers, rows, worksheet.title
    raise ValueError("仅支持 CSV/XLSX：%s" % path)


def resolve_file(root: Path, table: str, config: Dict[str, Any], required: bool) -> Optional[Path]:
    configured = config.get("files", {}).get(table)
    candidates = [configured] if configured else FILE_CANDIDATES[table]
    for candidate in candidates:
        if not candidate:
            continue
        path = (root / candidate).resolve()
        if root != path and root not in path.parents:
            raise ValueError("输入文件不得跳出 input-dir：%s" % candidate)
        if path.is_file():
            return path
    if required:
        raise FileNotFoundError("缺少必需输入 %s；候选文件：%s" % (table, ", ".join(FILE_CANDIDATES[table])))
    return None


def resolve_mapping(table: str, headers: Sequence[str], explicit: Dict[str, Any]) -> Dict[str, str]:
    normalized: Dict[str, List[str]] = defaultdict(list)
    for header in headers:
        normalized[header_key(header)].append(header)
    mapping = {}
    for canonical, aliases in SCHEMAS[table]["fields"].items():
        if canonical in explicit:
            source = text(explicit[canonical])
            if source not in headers:
                raise ValueError("%s 映射 %s -> %s 不存在" % (table, canonical, source))
            mapping[canonical] = source
            continue
        matches = []
        for alias in aliases:
            matches.extend(normalized.get(header_key(alias), []))
        matches = list(dict.fromkeys(matches))
        if len(matches) == 1:
            mapping[canonical] = matches[0]
    missing = [field for field in SCHEMAS[table]["required"] if field not in mapping]
    if missing:
        raise ValueError("%s 缺核心字段映射：%s" % (table, ", ".join(missing)))
    return mapping


def normalize_table(
    table: str, source: Sequence[Dict[str, Any]], mapping: Dict[str, str], default_currency: str, source_hash: str
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    rows, bad = [], []
    fields = list(SCHEMAS[table]["fields"].keys())
    for item in source:
        row = {field: text(item.get(mapping.get(field, ""), "")) for field in fields}
        row.update({"_source_file": item["_source_file"], "_source_row": item["_source_row"], "_source_sheet": item["_source_sheet"], "_source_hash": source_hash})
        reasons = []
        for required in SCHEMAS[table]["required"]:
            if not row.get(required):
                reasons.append("%s 为空" % required)
        for field in NUMBER_FIELDS.intersection(fields):
            if not row.get(field) and field not in SCHEMAS[table]["required"]:
                row[field] = None
                continue
            parsed = number(row.get(field))
            if parsed is None:
                if field in SCHEMAS[table]["required"]:
                    reasons.append("%s 无法解析" % field)
                row[field] = None
            else:
                row[field] = round(parsed, 6)
        for field in DATE_FIELDS.intersection(fields):
            if not row.get(field):
                row[field] = None
                continue
            parsed = date_value(row[field])
            if parsed is None and field in SCHEMAS[table]["required"]:
                reasons.append("%s 无法解析" % field)
            row[field] = parsed
        if "currency" in row:
            row["currency"] = (row["currency"] or default_currency).upper()
        if reasons:
            bad.append({"table": table, "source_file": item["_source_file"], "source_row": item["_source_row"], "reasons": "；".join(reasons), "raw_record": json.dumps({k: item.get(v) for k, v in mapping.items()}, ensure_ascii=False, default=str)})
        else:
            rows.append(row)
    return rows, bad


def risk_priority(score: int) -> str:
    return "critical" if score >= 6 else "high" if score >= 4 else "medium" if score >= 2 else "low"


class Builder:
    def __init__(self) -> None:
        self.findings: List[Dict[str, Any]] = []
        self.evidence: List[Dict[str, Any]] = []
        self.cache: Dict[Tuple[str, Any, str], str] = {}

    def row_evidence(self, row: Dict[str, Any], field: str) -> str:
        key = (row["_source_file"], row["_source_row"], field)
        if key in self.cache:
            return self.cache[key]
        evidence_id = "EV-%06d" % (len(self.evidence) + 1)
        self.evidence.append({
            "evidence_id": evidence_id, "source_file": row["_source_file"],
            "source_hash": "sha256:" + row["_source_hash"], "sheet": row.get("_source_sheet"),
            "row": row["_source_row"], "field": field, "value": row.get(field, ""),
            "extraction_method": "normalized", "confidence": 1.0,
        })
        self.cache[key] = evidence_id
        return evidence_id

    def document_evidence(self, path: Path, snippet: str) -> str:
        key = (path.name, None, "document_text")
        if key in self.cache:
            return self.cache[key]
        evidence_id = "EV-%06d" % (len(self.evidence) + 1)
        self.evidence.append({
            "evidence_id": evidence_id, "source_file": path.name,
            "source_hash": "sha256:" + sha256_file(path), "sheet": None, "row": None,
            "field": "document_text", "value": snippet, "extraction_method": "direct_text",
            "confidence": 1.0,
        })
        self.cache[key] = evidence_id
        return evidence_id

    def add(self, finding_type: str, title: str, score: int, strength: str,
            records: Sequence[Dict[str, Any]], fields: Sequence[str], facts: Sequence[str],
            inferences: Sequence[str], questions: Sequence[str], next_steps: Sequence[str],
            factors: Sequence[Dict[str, Any]], extra_refs: Sequence[str] = ()) -> Dict[str, Any]:
        refs = list(extra_refs)
        for record in records:
            refs.extend(self.row_evidence(record, field) for field in fields if field in record)
        entities, seen = [], set()
        candidates = (("vendor", "vendor_id"), ("vendor", "bidder_id"), ("employee", "employee_id"), ("buyer", "buyer_id"), ("lot", "lot_id"))
        for record in records:
            for entity_type, field in candidates:
                value = record.get(field)
                key = (entity_type, value)
                if value and key not in seen:
                    entities.append({"type": entity_type, "id": value})
                    seen.add(key)
        finding = {
            "finding_id": "PROC-%06d" % (len(self.findings) + 1), "skill": SKILL,
            "finding_type": finding_type, "title": title, "risk_priority": risk_priority(score),
            "evidence_strength": strength, "risk_score": score, "risk_factors": list(factors),
            "entities": entities, "facts": list(facts), "inferences": list(inferences),
            "hypotheses": [], "open_questions": list(questions), "evidence_refs": list(dict.fromkeys(refs)),
            "recommended_next_steps": list(next_steps), "human_review_required": True,
        }
        self.findings.append(finding)
        return finding


def whitelist_keys(config: Dict[str, Any], attribute: str) -> set:
    return {match_key(value) for value in config.get("whitelists", {}).get(attribute, []) if match_key(value)}


def add_shared_attribute_findings(vendors: List[Dict[str, Any]], employees: List[Dict[str, Any]], config: Dict[str, Any], builder: Builder, edges: List[Dict[str, Any]]) -> None:
    settings = {
        "bank_account": (4, "strong"), "phone": (3, "strong"), "email": (3, "strong"),
        "address": (1, "weak"), "legal_representative": (2, "moderate"),
    }
    for attribute, (score, strength) in settings.items():
        groups: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        allowed = whitelist_keys(config, attribute)
        for vendor in vendors:
            key = match_key(vendor.get(attribute))
            if key and key not in allowed:
                groups[key].append(vendor)
        for group in groups.values():
            if len(group) < 2:
                continue
            kind = attribute.replace("_", "-")
            finding = builder.add("shared-" + kind, "多个供应商共享%s" % attribute, score, strength, group,
                                  ("vendor_id", attribute),
                                  ["%d 个供应商的归一化 %s 相同" % (len(group), attribute)],
                                  ["该共享属性需要核对主体独立性和主数据准确性"],
                                  ["是否为集团公共信息、园区地址、平台账户或录入错误？"],
                                  ["核对开户资料、工商资料和供应商准入文件"],
                                  [{"factor": "shared_" + attribute, "points": score}])
            for index, left in enumerate(group):
                for right in group[index + 1 :]:
                    edges.append({"source": "vendor:" + left["vendor_id"], "target": "vendor:" + right["vendor_id"], "type": "shares_" + attribute, "evidence_refs": finding["evidence_refs"]})
    employee_settings = {"bank_account": (4, "strong"), "phone": (3, "strong"), "email": (3, "strong"), "address": (1, "weak")}
    for attribute, (score, strength) in employee_settings.items():
        allowed = whitelist_keys(config, attribute)
        by_value: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for vendor in vendors:
            key = match_key(vendor.get(attribute))
            if key and key not in allowed:
                by_value[key].append(vendor)
        for employee in employees:
            key = match_key(employee.get(attribute))
            if not key or key in allowed:
                continue
            for vendor in by_value.get(key, []):
                finding = builder.add("employee-vendor-shared-" + attribute.replace("_", "-"), "员工与供应商共享%s" % attribute,
                                      score, strength, (employee, vendor), ("employee_id", "vendor_id", attribute),
                                      ["员工 %s 与供应商 %s 的归一化 %s 相同" % (employee["employee_id"], vendor["vendor_id"], attribute)],
                                      ["该属性匹配构成潜在未披露关联的复核线索"],
                                      ["是否为公共信息、数据录入错误或已申报关系？"],
                                      ["核对利益冲突申报、员工资料和供应商准入材料"],
                                      [{"factor": "employee_vendor_shared_" + attribute, "points": score}])
                edges.append({"source": "employee:" + employee["employee_id"], "target": "vendor:" + vendor["vendor_id"], "type": "shares_" + attribute, "evidence_refs": finding["evidence_refs"]})


def add_price_findings(pos: List[Dict[str, Any]], config: Dict[str, Any], builder: Builder) -> None:
    minimum = int(config.get("price_outlier_min_group_size", 5))
    threshold = float(config.get("price_outlier_robust_z", 3.5))
    groups: Dict[Tuple[str, str, str, str], List[Dict[str, Any]]] = defaultdict(list)
    for po in pos:
        subject = match_key(po.get("item")) or ("category:" + match_key(po.get("category")))
        groups[(subject, match_key(po.get("unit")), match_key(po.get("region")), po.get("currency", ""))].append(po)
    for key, group in groups.items():
        if len(group) < minimum:
            continue
        values = [po["unit_price"] for po in group]
        median = statistics.median(values)
        mad = statistics.median(abs(value - median) for value in values)
        if mad == 0:
            continue
        for po in group:
            robust_z = 0.6745 * (po["unit_price"] - median) / mad
            if robust_z > threshold:
                builder.add("price-outlier", "同类采购单价显著高于 peer group", 2, "moderate", (po,),
                            ("po_id", "vendor_id", "item", "unit", "region", "unit_price"),
                            ["peer group 中位单价 %.2f、MAD %.2f，本单 robust z-score %.2f" % (median, mad, robust_z)],
                            ["该单价是同类组高额统计离群点"],
                            ["规格、税、运费、质量、交期和采购时间是否可比？"],
                            ["补齐规格与报价依据，重新确认 peer group 后复核"],
                            [{"factor": "price_outlier", "points": 2, "robust_z": round(robust_z, 4), "peer_group": list(key)}])


def add_split_findings(pos: List[Dict[str, Any]], config: Dict[str, Any], builder: Builder) -> None:
    window_days = int(config.get("split_window_days", 7))
    for threshold in config.get("approval_thresholds", []):
        limit = number(threshold.get("amount"))
        currency = text(threshold.get("currency")).upper()
        if limit is None:
            continue
        groups: Dict[Tuple[str, str, str, str], List[Dict[str, Any]]] = defaultdict(list)
        for po in pos:
            if currency and po["currency"] != currency:
                continue
            groups[(po["buyer_id"], po["vendor_id"], match_key(po.get("category")), po["currency"])].append(po)
        for group in groups.values():
            group.sort(key=lambda po: po["order_date"])
            seen = set()
            for index, first in enumerate(group):
                first_date = date.fromisoformat(first["order_date"])
                window = [po for po in group[index:] if (date.fromisoformat(po["order_date"]) - first_date).days <= window_days]
                key = tuple(sorted(po["po_id"] for po in window))
                if len(window) >= 2 and key not in seen and all(po["total_amount"] < limit for po in window) and sum(po["total_amount"] for po in window) > limit:
                    seen.add(key)
                    total = sum(po["total_amount"] for po in window)
                    builder.add("split-order", "短期多笔采购合计超过审批阈值", 3, "moderate", window,
                                ("po_id", "buyer_id", "vendor_id", "category", "total_amount", "order_date"),
                                ["%d 笔单笔低于 %.2f 的订单合计 %.2f" % (len(window), limit, total)],
                                ["该组合符合配置的疑似拆单条件"],
                                ["是否为框架协议分批交付、独立需求或系统拆行？"],
                                ["核对需求申请、审批层级、合同和交付计划"],
                                [{"factor": "split_order", "points": 3, "rule_id": threshold.get("rule_id"), "window_days": window_days}])


def add_process_findings(pos: List[Dict[str, Any]], payments: List[Dict[str, Any]], builder: Builder) -> None:
    po_by_id = {po["po_id"]: po for po in pos}
    for po in pos:
        if po.get("approval_date") and po["order_date"] < po["approval_date"]:
            builder.add("process-order-before-approval", "采购订单日期早于审批完成日期", 2, "strong", (po,),
                        ("po_id", "order_date", "approval_date"), ["订单日期 %s 早于审批日期 %s" % (po["order_date"], po["approval_date"])],
                        ["流程时间顺序需要核对"], ["是否为紧急采购、系统补录或日期口径差异？"],
                        ["调取审批日志和订单创建时间戳"], [{"factor": "order_before_approval", "points": 2}])
        if po.get("receipt_date") and po["receipt_date"] < po["order_date"]:
            builder.add("process-receipt-before-order", "收货日期早于采购订单日期", 2, "strong", (po,),
                        ("po_id", "receipt_date", "order_date"), ["收货日期 %s 早于订单日期 %s" % (po["receipt_date"], po["order_date"])],
                        ["可能存在先执行后补单或日期数据错误"], ["是否为历史补录、退换货或接口口径差异？"],
                        ["核对收货单、系统日志和合同生效时间"], [{"factor": "receipt_before_order", "points": 2}])
    for payment in payments:
        po = po_by_id.get(payment["po_id"])
        if po and po.get("approval_date") and payment["payment_date"] < po["approval_date"]:
            builder.add("process-payment-before-approval", "付款日期早于采购审批完成日期", 3, "strong", (payment, po),
                        ("payment_id", "po_id", "payment_date", "approval_date", "amount"),
                        ["付款日期 %s 早于关联 PO 审批日期 %s" % (payment["payment_date"], po["approval_date"])],
                        ["付款流程顺序需要复核"], ["是否为预付款、紧急授权、补录或日期口径差异？"],
                        ["核对付款审批、银行流水和合同预付款条款"], [{"factor": "payment_before_approval", "points": 3}])


def add_concentration_findings(pos: List[Dict[str, Any]], config: Dict[str, Any], builder: Builder) -> None:
    share_limit = float(config.get("concentration_share", 0.7))
    minimum_orders = int(config.get("concentration_min_orders", 3))
    groups: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for po in pos:
        groups[po["buyer_id"]].append(po)
    for buyer, group in groups.items():
        total = sum(po["total_amount"] for po in group)
        by_vendor: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for po in group:
            by_vendor[po["vendor_id"]].append(po)
        for vendor, subset in by_vendor.items():
            share = sum(po["total_amount"] for po in subset) / total if total else 0
            if len(subset) >= minimum_orders and share >= share_limit:
                builder.add("buyer-vendor-concentration", "采购员对单一供应商的金额集中度较高", 1, "moderate", subset,
                            ("po_id", "buyer_id", "vendor_id", "total_amount"),
                            ["采购员 %s 的 %d 笔订单中，供应商 %s 金额占比 %.2f%%" % (buyer, len(group), vendor, share * 100)],
                            ["该集中度达到配置的复核阈值"], ["是否存在独家技术、框架协议或地区供给限制？"],
                            ["查看供应商选择依据、市场询价和合同覆盖范围"],
                            [{"factor": "buyer_vendor_concentration", "points": 1, "share": round(share, 6)}])


def ngrams(value: str, minimum: int, maximum: int) -> Counter:
    compact = re.sub(r"\s+", "", value)
    result: Counter = Counter()
    for size in range(minimum, maximum + 1):
        result.update(compact[index:index + size] for index in range(max(len(compact) - size + 1, 0)))
    return result


def tfidf_vectors(documents: Sequence[str], minimum: int, maximum: int) -> List[Dict[str, float]]:
    counts = [ngrams(document, minimum, maximum) for document in documents]
    document_frequency = Counter()
    for count in counts:
        document_frequency.update(count.keys())
    total = len(documents)
    vectors = []
    for count in counts:
        length = sum(count.values()) or 1
        vectors.append({term: (frequency / length) * (math.log((1 + total) / (1 + document_frequency[term])) + 1) for term, frequency in count.items()})
    return vectors


def cosine(left: Dict[str, float], right: Dict[str, float]) -> float:
    dot = sum(value * right.get(term, 0.0) for term, value in left.items())
    left_norm = math.sqrt(sum(value * value for value in left.values()))
    right_norm = math.sqrt(sum(value * value for value in right.values()))
    return dot / (left_norm * right_norm) if left_norm and right_norm else 0.0


def safe_document(root: Path, relative: str) -> Path:
    path = (root / relative).resolve()
    if root != path and root not in path.parents:
        raise ValueError("document_path 不得跳出 input-dir：%s" % relative)
    if path.suffix.lower() not in (".txt", ".md"):
        raise ValueError("核心文本比较仅支持 .txt/.md：%s" % relative)
    return path


def add_bid_findings(root: Path, bids: List[Dict[str, Any]], config: Dict[str, Any], builder: Builder) -> List[str]:
    warnings = []
    minimum, maximum = int(config.get("bid_ngram_min", 2)), int(config.get("bid_ngram_max", 5))
    threshold = float(config.get("bid_similarity_threshold", 0.78))
    templates = [text(value) for value in config.get("common_template_phrases", []) if text(value)]
    docs, doc_bids, paths = [], [], []
    for bid in bids:
        if not bid.get("document_path"):
            continue
        try:
            path = safe_document(root, bid["document_path"])
            raw = path.read_text(encoding="utf-8-sig")
        except Exception as exc:
            warnings.append("投标文本跳过 %s：%s" % (bid.get("document_path"), exc))
            continue
        normalized = unicodedata.normalize("NFKC", raw).lower()
        for phrase in templates:
            normalized = normalized.replace(phrase.lower(), "")
        if len(re.sub(r"\s+", "", normalized)) < 20:
            warnings.append("投标文本过短：%s" % bid["document_path"])
            continue
        docs.append(normalized)
        doc_bids.append(bid)
        paths.append(path)
    if docs:
        vectors = tfidf_vectors(docs, minimum, maximum)
        for left_index, left in enumerate(doc_bids):
            for right_index in range(left_index + 1, len(doc_bids)):
                right = doc_bids[right_index]
                if left["lot_id"] != right["lot_id"]:
                    continue
                score = cosine(vectors[left_index], vectors[right_index])
                if score >= threshold:
                    refs = [builder.document_evidence(paths[left_index], docs[left_index][:180]), builder.document_evidence(paths[right_index], docs[right_index][:180])]
                    builder.add("bid-text-similarity", "同一标段投标文本高度相似，需回到原文复核", 2, "moderate", (left, right),
                                ("tender_id", "lot_id", "bidder_id", "document_path"),
                                ["字符级 TF-IDF 余弦相似度 %.6f，高于本次配置阈值 %.6f" % (score, threshold)],
                                ["该文件对构成高相似复核线索，不等于串标认定"],
                                ["相同内容是招标公共模板还是投标人自行撰写？"],
                                ["定位相似片段，并核对文件元数据、报价、联系人和保证金来源"],
                                [{"factor": "bid_text_similarity", "points": 2, "similarity": round(score, 6), "threshold": threshold, "ngram_range": [minimum, maximum]}], refs)
    close_ratio = float(config.get("bid_price_close_ratio", 0.01))
    lots: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for bid in bids:
        lots[bid["lot_id"]].append(bid)
    for lot, group in lots.items():
        prices = [bid["bid_price"] for bid in group]
        if len(prices) >= 3:
            median = statistics.median(prices)
            spread = (max(prices) - min(prices)) / median if median else 0
            if spread <= close_ratio:
                builder.add("bid-price-pattern", "同一标段多家报价异常接近", 1, "moderate", group,
                            ("tender_id", "lot_id", "bidder_id", "bid_price"),
                            ["标段 %s 的 %d 个报价极差/中位数为 %.6f" % (lot, len(group), spread)],
                            ["报价接近是弱红旗，需要结合计价规则解释"],
                            ["是否存在统一预算、限价、税率或标准计价公式？"],
                            ["核对报价明细、编制痕迹和提交信息"],
                            [{"factor": "bid_price_close", "points": 1, "spread_ratio": round(spread, 6)}])
    return warnings


def write_csv(path: Path, rows: Sequence[Dict[str, Any]], fields: Sequence[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fields), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_jsonl(path: Path, rows: Sequence[Dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def validate(builder: Builder) -> None:
    evidence_ids = {item["evidence_id"] for item in builder.evidence}
    for finding in builder.findings:
        if not finding["evidence_refs"] or not set(finding["evidence_refs"]).issubset(evidence_ids):
            raise ValueError("finding %s 的 evidence 引用不完整" % finding["finding_id"])
        if not finding["human_review_required"]:
            raise ValueError("finding 必须要求人工复核")


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path)
    parser.add_argument("--config", type=Path, required=True, help="必传；含审批阈值、相似度阈值、白名单等参数")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--check-env", action="store_true")
    args = parser.parse_args(argv)
    if args.check_env:
        try:
            import openpyxl
            xlsx = "available (%s)" % openpyxl.__version__
        except ImportError:
            xlsx = "unavailable; CSV/TXT mode still works"
        print(json.dumps({"python": sys.version.split()[0], "recommended": ">=3.10", "xlsx_support": xlsx, "network_required": False}, ensure_ascii=False))
        return 0
    if not args.input_dir or not args.output:
        parser.error("正式运行需要 --input-dir 和 --output")
    root = args.input_dir.resolve()
    if not root.is_dir():
        raise NotADirectoryError(root)
    output = args.output.resolve()
    if output.exists() and any(output.iterdir()):
        raise FileExistsError("输出目录非空；请使用新目录：%s" % output)
    output.mkdir(parents=True, exist_ok=True)
    config = load_json(args.config.resolve() if args.config else None)
    default_currency = text(config.get("default_currency", "CNY")).upper()
    started = datetime.now(timezone.utc)
    tables: Dict[str, List[Dict[str, Any]]] = {}
    mappings, input_files, quality, bad_rows = {}, [], {}, []
    for table in SCHEMAS:
        path = resolve_file(root, table, config, table in ("vendors", "purchase_orders"))
        if path is None:
            tables[table] = []
            quality[table] = {"status": "not_provided", "rows": 0, "valid_rows": 0, "bad_rows": 0}
            continue
        digest = sha256_file(path)
        headers, source_rows, sheet = read_table(path)
        mapping = resolve_mapping(table, headers, config.get("field_maps", {}).get(table, {}))
        rows, bad = normalize_table(table, source_rows, mapping, default_currency, digest)
        tables[table] = rows
        mappings[table] = mapping
        bad_rows.extend(bad)
        input_files.append({"path": path.name, "sha256": digest, "sheet": sheet, "table": table})
        quality[table] = {"status": "loaded", "rows": len(source_rows), "valid_rows": len(rows), "bad_rows": len(bad), "mapping": mapping}

    builder, edges = Builder(), []
    add_shared_attribute_findings(tables["vendors"], tables["employees"], config, builder, edges)
    add_price_findings(tables["purchase_orders"], config, builder)
    add_split_findings(tables["purchase_orders"], config, builder)
    add_process_findings(tables["purchase_orders"], tables["payments"], builder)
    add_concentration_findings(tables["purchase_orders"], config, builder)
    warnings = add_bid_findings(root, tables["bids"], config, builder)

    for po in tables["purchase_orders"]:
        refs = [builder.row_evidence(po, "po_id"), builder.row_evidence(po, "total_amount")]
        edges.append({"source": "buyer:" + po["buyer_id"], "target": "vendor:" + po["vendor_id"], "type": "purchased_from", "amount": po["total_amount"], "currency": po["currency"], "evidence_refs": refs})
    nodes, node_seen = [], set()
    for node_type, table, field in (("vendor", "vendors", "vendor_id"), ("employee", "employees", "employee_id"), ("buyer", "purchase_orders", "buyer_id")):
        for row in tables[table]:
            node_id = node_type + ":" + row[field]
            if node_id not in node_seen:
                nodes.append({"id": node_id, "type": node_type, "label": row[field]})
                node_seen.add(node_id)
    validate(builder)

    for table, rows in tables.items():
        if rows:
            fields = list(SCHEMAS[table]["fields"].keys()) + ["_source_file", "_source_row", "_source_sheet"]
            write_csv(output / ("normalized_%s.csv" % table), rows, fields)
    write_csv(output / "bad_rows.csv", bad_rows, ("table", "source_file", "source_row", "reasons", "raw_record"))
    flat = [{
        "finding_id": f["finding_id"], "finding_type": f["finding_type"], "title": f["title"],
        "risk_priority": f["risk_priority"], "evidence_strength": f["evidence_strength"],
        "risk_score": f["risk_score"], "entities": json.dumps(f["entities"], ensure_ascii=False),
        "evidence_refs": "|".join(f["evidence_refs"]),
    } for f in builder.findings]
    write_csv(output / "findings.csv", flat, ("finding_id", "finding_type", "title", "risk_priority", "evidence_strength", "risk_score", "entities", "evidence_refs"))
    write_jsonl(output / "findings.jsonl", builder.findings)
    write_jsonl(output / "evidence.jsonl", builder.evidence)
    (output / "relationship_graph.json").write_text(json.dumps({"schema_version": VERSION, "nodes": nodes, "edges": edges}, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")

    entity_findings: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for finding in builder.findings:
        for entity in finding["entities"]:
            if entity["type"] == "vendor":
                entity_findings[entity["id"]].append(finding)
    candidates = []
    for vendor, findings in entity_findings.items():
        types = {f["finding_type"] for f in findings}
        if len(types) >= 2 and any(f["risk_priority"] in ("high", "critical") for f in findings):
            candidates.append((vendor, findings))
    if candidates:
        vendor, findings = max(candidates, key=lambda item: (len({f["finding_type"] for f in item[1]}), len(item[1])))
        handoff = {
            "schema_version": VERSION, "status": "recommended", "source_skill": SKILL,
            "finding_ids": [f["finding_id"] for f in findings], "entities": [{"type": "vendor", "id": vendor}],
            "reason_for_handoff": "多个独立采购红旗指向同一供应商，建议由有权人员决定是否进入调查辅助流程",
            "known_facts": list(dict.fromkeys(fact for f in findings for fact in f["facts"]))[:20],
            "hypotheses": ["这些红旗可能有共同原因，也可能由多个独立且合理的业务情形造成"],
            "recommended_scope": ["仅围绕列示 finding 和关联期间核验主数据、审批、报价、付款与利益冲突申报"],
            "human_approval_required": True,
        }
    else:
        handoff = {"schema_version": VERSION, "status": "not_recommended", "source_skill": SKILL, "finding_ids": [], "entities": [], "reason_for_handoff": "未达到多类独立红旗的建议条件", "known_facts": [], "hypotheses": [], "recommended_scope": [], "human_approval_required": True}
    (output / "investigation_handoff.json").write_text(json.dumps(handoff, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")

    skipped = [table + " 模块：未提供输入" for table in ("employees", "payments", "bids") if not tables[table]]
    if not config.get("approval_thresholds"):
        skipped.append("split-order：未提供 approval_thresholds")
    quality_lines = ["# 采购数据质量报告", "", "- Bad rows 合计：%d" % len(bad_rows), "", "## 各表", ""]
    quality_lines.extend("- `%s`: `%s`" % (table, json.dumps(value, ensure_ascii=False)) for table, value in quality.items())
    if warnings or skipped:
        quality_lines.extend(["", "## 警告与跳过", ""] + ["- " + item for item in warnings + skipped])
    (output / "data_quality.md").write_text("\n".join(quality_lines) + "\n", encoding="utf-8")

    counts, priorities = Counter(f["finding_type"] for f in builder.findings), Counter(f["risk_priority"] for f in builder.findings)
    summary = [
        "# 采购红旗确定性摘要", "", "- Findings：%d；Evidence：%d。" % (len(builder.findings), len(builder.evidence)),
        "- 风险优先级：`%s`" % json.dumps(dict(priorities), ensure_ascii=False, sort_keys=True),
        "- 类型：`%s`" % json.dumps(dict(counts), ensure_ascii=False, sort_keys=True),
        "- 调查移交建议：`%s`。" % handoff["status"], "",
        "> 共享属性、价格离群、流程异常和文本相似均为复核线索。它们不能单独或自动证明串标、利益输送或舞弊。", "",
        "## 建议复核顺序", "", "1. 先核验共享银行账号和员工—供应商强属性是否准确及已申报。",
        "2. 再回到 PO、审批、付款和投标原文检查同一主体上的多模块组合。",
        "3. 主动核对公共地址、模板、独家供应、紧急采购和系统补录等替代解释。",
    ]
    (output / "summary.md").write_text("\n".join(summary) + "\n", encoding="utf-8")

    script = Path(__file__).resolve()
    if args.config:
        input_files.append({"path": args.config.name, "sha256": sha256_file(args.config.resolve()), "table": "config"})
    manifest = {
        "run_id": "PROC-%s-%s" % (started.strftime("%Y%m%dT%H%M%SZ"), hashlib.sha256(json.dumps(input_files, sort_keys=True).encode()).hexdigest()[:8]),
        "skill": SKILL, "skill_version": VERSION, "started_at": started.isoformat(), "finished_at": datetime.now(timezone.utc).isoformat(),
        "input_files": input_files, "field_mappings": mappings,
        "parameters": {key: config.get(key) for key in ("config_version", "default_currency", "price_outlier_min_group_size", "price_outlier_robust_z", "split_window_days", "approval_thresholds", "bid_similarity_threshold", "bid_ngram_min", "bid_ngram_max", "bid_price_close_ratio", "concentration_share", "concentration_min_orders", "common_template_phrases", "whitelists")},
        "scripts": {script.name: "sha256:" + sha256_file(script)}, "warnings": warnings, "skipped_modules": skipped,
        "network_access": False,
        "outputs": sorted(path.name for path in output.iterdir()) + ["run_manifest.json"],
    }
    (output / "run_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), "findings": len(builder.findings), "evidence": len(builder.evidence), "handoff": handoff["status"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print("ERROR: %s" % exc, file=sys.stderr)
        raise SystemExit(2)
