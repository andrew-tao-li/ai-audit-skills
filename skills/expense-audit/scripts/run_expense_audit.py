#!/usr/bin/env python3
"""Offline, deterministic expense red-flag scanner."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import statistics
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

VERSION = "0.1.1"
SKILL = "expense-audit"

ALIASES = {
    "expense_id": ["expense_id", "id", "claim_id", "report_id", "单据号", "报销单号", "费用编号"],
    "employee_id": ["employee_id", "emp_id", "employee", "staff_id", "员工编号", "工号", "员工"],
    "department": ["department", "dept", "部门"],
    "expense_type": ["expense_type", "category", "type", "费用类型", "费用类别", "报销类型"],
    "expense_date": ["expense_date", "date", "transaction_date", "发生日期", "费用日期", "报销日期"],
    "submit_date": ["submit_date", "submission_date", "提交日期"],
    "amount": ["amount", "expense_amount", "claim_amount", "金额", "报销金额", "费用金额"],
    "currency": ["currency", "币种", "货币"],
    "vendor_name": ["vendor_name", "vendor", "merchant", "supplier", "商户", "商家", "供应商"],
    "invoice_number": ["invoice_number", "invoice_no", "invoice", "发票号", "发票号码"],
    "invoice_date": ["invoice_date", "发票日期"],
    "project_code": ["project_code", "project", "项目编号", "项目代码"],
    "approver": ["approver", "approved_by", "审批人"],
    "payment_date": ["payment_date", "付款日期", "支付日期"],
    "business_purpose": ["business_purpose", "purpose", "description", "事由", "用途", "业务目的"],
}
CORE_FIELDS = ("expense_id", "employee_id", "amount", "expense_date")
OUTPUT_FIELDS = tuple(ALIASES.keys())


def norm_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(value))).strip()


def norm_header(value: Any) -> str:
    return re.sub(r"[\s_\-./\\]+", "", norm_text(value).lower())


def match_key(value: Any) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", norm_text(value).lower())


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_amount(value: Any) -> Optional[float]:
    text = norm_text(value)
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"[^0-9.+-]", "", text.replace(",", ""))
    try:
        number = float(cleaned)
        return -number if negative else number
    except (TypeError, ValueError):
        return None


def parse_date(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = norm_text(value)
    if not text:
        return None
    text = text.split("T")[0].split(" ")[0]
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def read_table(path: Path, sheet: Optional[str]) -> Tuple[List[str], List[Dict[str, Any]], Optional[str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            headers = [norm_text(h) for h in (reader.fieldnames or [])]
            rows = []
            for row_number, raw in enumerate(reader, start=2):
                row = {norm_text(k): v for k, v in raw.items() if k is not None}
                row["_source_row"] = row_number
                row["_source_sheet"] = None
                rows.append(row)
            return headers, rows, None
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("读取 XLSX 需要 openpyxl；请离线安装或另存为 UTF-8 CSV") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        if sheet and sheet not in workbook.sheetnames:
            raise ValueError("找不到工作表：%s；可用工作表：%s" % (sheet, ", ".join(workbook.sheetnames)))
        worksheet = workbook[sheet] if sheet else workbook[workbook.sheetnames[0]]
        values = worksheet.iter_rows(values_only=True)
        first = next(values, None)
        if first is None:
            return [], [], worksheet.title
        headers = [norm_text(v) for v in first]
        rows = []
        for row_number, values_row in enumerate(values, start=2):
            row = {headers[index]: value for index, value in enumerate(values_row) if index < len(headers)}
            row["_source_row"] = row_number
            row["_source_sheet"] = worksheet.title
            rows.append(row)
        workbook.close()
        return headers, rows, worksheet.title
    raise ValueError("仅支持 .csv 和 .xlsx 输入")


def load_json(path: Optional[Path]) -> Dict[str, Any]:
    if path is None:
        return {}
    with path.open("r", encoding="utf-8-sig") as handle:
        value = json.load(handle)
    if not isinstance(value, dict):
        raise ValueError("JSON 配置顶层必须是对象")
    return value


def resolve_mapping(headers: Sequence[str], explicit: Dict[str, Any]) -> Tuple[Dict[str, str], List[str]]:
    by_normalized = defaultdict(list)
    for header in headers:
        by_normalized[norm_header(header)].append(header)
    mapping: Dict[str, str] = {}
    warnings: List[str] = []
    for canonical in OUTPUT_FIELDS:
        if canonical in explicit:
            source = norm_text(explicit[canonical])
            if source not in headers:
                raise ValueError("字段映射 %s -> %s 不存在于输入表" % (canonical, source))
            mapping[canonical] = source
            continue
        matches: List[str] = []
        for alias in ALIASES[canonical]:
            matches.extend(by_normalized.get(norm_header(alias), []))
        matches = list(dict.fromkeys(matches))
        if len(matches) == 1:
            mapping[canonical] = matches[0]
        elif len(matches) > 1:
            warnings.append("字段 %s 存在多个候选：%s" % (canonical, ", ".join(matches)))
    missing = [field for field in CORE_FIELDS if field not in mapping]
    if missing:
        raise ValueError("缺少核心字段映射：%s。请使用 --field-map 提供明确映射" % ", ".join(missing))
    return mapping, warnings


def normalize_rows(
    source_rows: Sequence[Dict[str, Any]], mapping: Dict[str, str], default_currency: str
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], Dict[str, int]]:
    clean: List[Dict[str, Any]] = []
    bad: List[Dict[str, Any]] = []
    changes = {"trimmed_or_normalized_text_values": 0, "parsed_amount_values": 0, "parsed_date_values": 0}
    for source in source_rows:
        row: Dict[str, Any] = {"_source_row": source["_source_row"], "_source_sheet": source["_source_sheet"]}
        for canonical in OUTPUT_FIELDS:
            raw = source.get(mapping.get(canonical, ""), "")
            normalized = norm_text(raw)
            row[canonical] = normalized
            if normalized != ("" if raw is None else str(raw).strip()):
                changes["trimmed_or_normalized_text_values"] += 1
        amount = parse_amount(row["amount"])
        expense_date = parse_date(row["expense_date"])
        reasons = []
        if not row["expense_id"]:
            reasons.append("expense_id 为空")
        if not row["employee_id"]:
            reasons.append("employee_id 为空")
        if amount is None or not math.isfinite(amount):
            reasons.append("amount 无法解析为有限数值")
        if expense_date is None:
            reasons.append("expense_date 无法解析")
        if reasons:
            bad.append({"source_row": source["_source_row"], "reasons": "；".join(reasons), "raw_record": json.dumps({k: source.get(k) for k in mapping.values()}, ensure_ascii=False, default=str)})
            continue
        row["amount"] = round(float(amount), 6)
        row["expense_date"] = expense_date
        row["currency"] = (row["currency"] or default_currency).upper()
        for date_field in ("submit_date", "invoice_date", "payment_date"):
            if row[date_field]:
                parsed = parse_date(row[date_field])
                row[date_field] = parsed or row[date_field]
        changes["parsed_amount_values"] += 1
        changes["parsed_date_values"] += 1
        clean.append(row)
    return clean, bad, changes


def priority(score: int) -> str:
    if score >= 6:
        return "critical"
    if score >= 4:
        return "high"
    if score >= 2:
        return "medium"
    return "low"


class ResultBuilder:
    def __init__(self, source_file: str, source_hash: str):
        self.source_file = source_file
        self.source_hash = "sha256:" + source_hash
        self.evidence: List[Dict[str, Any]] = []
        self.findings: List[Dict[str, Any]] = []
        self._evidence_cache: Dict[Tuple[int, str], str] = {}

    def evidence_for(self, record: Dict[str, Any], field: str) -> str:
        key = (int(record["_source_row"]), field)
        if key in self._evidence_cache:
            return self._evidence_cache[key]
        evidence_id = "EV-%06d" % (len(self.evidence) + 1)
        item = {
            "evidence_id": evidence_id,
            "source_file": self.source_file,
            "source_hash": self.source_hash,
            "sheet": record.get("_source_sheet"),
            "row": record["_source_row"],
            "field": field,
            "value": record.get(field, ""),
            "extraction_method": "normalized",
            "confidence": 1.0,
        }
        self.evidence.append(item)
        self._evidence_cache[key] = evidence_id
        return evidence_id

    def add(
        self,
        finding_type: str,
        title: str,
        score: int,
        strength: str,
        records: Sequence[Dict[str, Any]],
        evidence_fields: Sequence[str],
        facts: Sequence[str],
        inferences: Sequence[str],
        questions: Sequence[str],
        next_steps: Sequence[str],
        factors: Sequence[Dict[str, Any]],
    ) -> None:
        refs = []
        for record in records:
            for field in evidence_fields:
                refs.append(self.evidence_for(record, field))
        entities = []
        seen = set()
        for record in records:
            for entity_type, field in (("employee", "employee_id"), ("vendor", "vendor_name")):
                value = record.get(field)
                key = (entity_type, value)
                if value and key not in seen:
                    entities.append({"type": entity_type, "id": value})
                    seen.add(key)
        self.findings.append({
            "finding_id": "EXP-%06d" % (len(self.findings) + 1),
            "skill": SKILL,
            "finding_type": finding_type,
            "title": title,
            "risk_priority": priority(score),
            "evidence_strength": strength,
            "risk_score": score,
            "risk_factors": list(factors),
            "entities": entities,
            "facts": list(facts),
            "inferences": list(inferences),
            "hypotheses": [],
            "open_questions": list(questions),
            "evidence_refs": list(dict.fromkeys(refs)),
            "recommended_next_steps": list(next_steps),
            "human_review_required": True,
        })


def group_records(rows: Sequence[Dict[str, Any]], fields: Sequence[str], require_nonempty: Sequence[str] = ()) -> Iterable[List[Dict[str, Any]]]:
    groups: Dict[Tuple[Any, ...], List[Dict[str, Any]]] = defaultdict(list)
    for row in rows:
        if any(not row.get(field) for field in require_nonempty):
            continue
        key = tuple(row.get(field) for field in fields)
        groups[key].append(row)
    return (group for group in groups.values() if len(group) > 1)


def run_rules(rows: List[Dict[str, Any]], policy: Dict[str, Any], builder: ResultBuilder) -> Tuple[List[str], Dict[str, Any]]:
    skipped: List[str] = []
    exact_pairs = set()
    for group in group_records(rows, ("invoice_number", "currency", "amount"), ("invoice_number",)):
        ids = [r["expense_id"] for r in group]
        exact_pairs.update(frozenset((a, b)) for i, a in enumerate(ids) for b in ids[i + 1 :])
        builder.add("exact-duplicate-invoice", "发票号与金额重复，需核对是否重复报销", 4, "strong", group,
                    ("expense_id", "invoice_number", "amount"),
                    ["%d 条记录具有相同的非空发票号、币种和金额" % len(group)],
                    ["该组合构成可能重复报销的高优先级复核线索"],
                    ["是否为冲销、退款、作废后重提或发票拆分？"],
                    ["核对发票影像、报销状态和付款流水"],
                    [{"factor": "same_invoice_and_amount", "points": 4}])
    for group in group_records(rows, ("employee_id", "expense_date", "currency", "amount")):
        ids = [r["expense_id"] for r in group]
        exact_pairs.update(frozenset((a, b)) for i, a in enumerate(ids) for b in ids[i + 1 :])
        builder.add("exact-duplicate-employee-date-amount", "同员工同日同金额记录重复", 3, "moderate", group,
                    ("expense_id", "employee_id", "expense_date", "amount", "vendor_name"),
                    ["%d 条记录的员工、日期、币种和金额相同" % len(group)],
                    ["可能是重复提交，也可能是同日不同业务"],
                    ["商户、业务事由和原始凭证是否不同？"],
                    ["逐笔核对事由、商户、发票和审批流"],
                    [{"factor": "same_employee_date_amount", "points": 3}])

    near_days = int(policy.get("near_duplicate_window_days", 7))
    near_tolerance = float(policy.get("near_duplicate_amount_tolerance", 0.02))
    ordered = sorted(rows, key=lambda r: (r["employee_id"], r["expense_date"], r["expense_id"]))
    near_seen = set()
    for index, left in enumerate(ordered):
        left_date = date.fromisoformat(left["expense_date"])
        for right in ordered[index + 1 :]:
            if right["employee_id"] != left["employee_id"]:
                if right["employee_id"] > left["employee_id"]:
                    break
                continue
            delta = abs((date.fromisoformat(right["expense_date"]) - left_date).days)
            if delta > near_days:
                continue
            pair = frozenset((left["expense_id"], right["expense_id"]))
            if pair in exact_pairs or pair in near_seen:
                continue
            if not match_key(left["vendor_name"]) or match_key(left["vendor_name"]) != match_key(right["vendor_name"]):
                continue
            if left["currency"] != right["currency"]:
                continue
            difference = abs(left["amount"] - right["amount"]) / max(abs(left["amount"]), abs(right["amount"]), 0.01)
            if difference <= near_tolerance:
                near_seen.add(pair)
                builder.add("near-duplicate", "同员工同商户短期内出现近似金额", 2, "moderate", (left, right),
                            ("expense_id", "employee_id", "expense_date", "vendor_name", "amount", "invoice_number"),
                            ["两条记录相隔 %d 天，金额差异比例 %.4f" % (delta, difference)],
                            ["该组合符合 possible near duplicate 的配置条件"],
                            ["是否为同一次消费的更正、补录或分次结算？"],
                            ["对照发票影像、消费时间和支付流水"],
                            [{"factor": "near_duplicate", "points": 2, "window_days": near_days, "tolerance": near_tolerance}])

    limits = policy.get("limits", [])
    if not limits:
        skipped.append("policy-threshold：未提供 limits")
    for row in rows:
        for limit in limits:
            if limit.get("expense_type") and match_key(limit["expense_type"]) != match_key(row["expense_type"]):
                continue
            if limit.get("currency") and str(limit["currency"]).upper() != row["currency"]:
                continue
            maximum = parse_amount(limit.get("max_amount"))
            if maximum is not None and row["amount"] > maximum:
                over = row["amount"] - maximum
                builder.add("policy-threshold", "费用金额超过适用制度上限", 3, "strong", (row,),
                            ("expense_id", "expense_type", "amount", "currency"),
                            ["金额 %.2f %s 超过配置上限 %.2f %s" % (row["amount"], row["currency"], maximum, row["currency"])],
                            ["该记录触发用户提供的制度阈值 %s" % limit.get("rule_id", "未命名规则")],
                            ["是否存在例外审批、多人费用或可报销附加项目？"],
                            ["核对制度版本、例外审批和费用明细"],
                            [{"factor": "policy_threshold", "points": 3, "rule_id": limit.get("rule_id"), "over_amount": round(over, 2)}])

    thresholds = policy.get("approval_thresholds", [])
    if not thresholds:
        skipped.append("split-expense：未提供 approval_thresholds")
    split_days = int(policy.get("split_window_days", 3))
    for threshold in thresholds:
        amount_threshold = parse_amount(threshold.get("amount"))
        currency = str(threshold.get("currency", "")).upper()
        if amount_threshold is None:
            continue
        groups: Dict[Tuple[str, str, str], List[Dict[str, Any]]] = defaultdict(list)
        for row in rows:
            if currency and row["currency"] != currency:
                continue
            groups[(row["employee_id"], match_key(row["vendor_name"]), row["currency"])].append(row)
        split_seen = set()
        for group in groups.values():
            group.sort(key=lambda r: r["expense_date"])
            for start, first in enumerate(group):
                end_date = date.fromisoformat(first["expense_date"])
                window = [r for r in group[start:] if (date.fromisoformat(r["expense_date"]) - end_date).days <= split_days]
                if len(window) < 2 or any(r["amount"] >= amount_threshold for r in window):
                    continue
                total = sum(r["amount"] for r in window)
                key = tuple(sorted(r["expense_id"] for r in window))
                if total > amount_threshold and key not in split_seen:
                    split_seen.add(key)
                    builder.add("split-expense", "短期多笔费用合计超过审批阈值", 3, "moderate", window,
                                ("expense_id", "employee_id", "expense_date", "vendor_name", "amount"),
                                ["%d 笔单笔低于 %.2f %s 的费用合计 %.2f %s" % (len(window), amount_threshold, window[0]["currency"], total, window[0]["currency"])],
                                ["该组合符合疑似拆单的配置条件"],
                                ["是否为独立业务、分期结算或不同费用承担人？"],
                                ["核对采购/服务内容、审批层级和原始支付时间"],
                                [{"factor": "split_expense", "points": 3, "rule_id": threshold.get("rule_id"), "window_days": split_days}])

    if policy.get("weekend_check", False):
        for row in rows:
            if date.fromisoformat(row["expense_date"]).weekday() >= 5:
                builder.add("weekend-signal", "费用发生在周末，需结合业务安排解释", 1, "weak", (row,),
                            ("expense_id", "expense_date", "business_purpose"),
                            ["费用日期 %s 为周末" % row["expense_date"]],
                            ["这是弱风险信号，不能单独说明费用不真实"],
                            ["是否有值班、出差、客户现场或跨时区安排？"],
                            ["核对日程、出差审批和业务事由"],
                            [{"factor": "weekend", "points": 1}])
    else:
        skipped.append("weekend-signal：weekend_check 未启用")

    min_group = int(policy.get("outlier_min_group_size", 5))
    z_threshold = float(policy.get("outlier_robust_z", 3.5))
    peer_groups: Dict[Tuple[str, str], List[Dict[str, Any]]] = defaultdict(list)
    for row in rows:
        peer_groups[(match_key(row["expense_type"]), row["currency"])].append(row)
    valid_groups = 0
    for key, group in peer_groups.items():
        if len(group) < min_group:
            continue
        values = [r["amount"] for r in group]
        median = statistics.median(values)
        mad = statistics.median([abs(value - median) for value in values])
        if mad == 0:
            continue
        valid_groups += 1
        for row in group:
            robust_z = 0.6745 * (row["amount"] - median) / mad
            if robust_z > z_threshold:
                builder.add("robust-outlier", "同类费用中的高额稳健统计离群点", 2, "moderate", (row,),
                            ("expense_id", "expense_type", "amount", "currency"),
                            ["同类组中位数 %.2f，MAD %.2f，本记录 robust z-score %.2f" % (median, mad, robust_z)],
                            ["该记录相对同类组明显偏高，仅用于复核排序"],
                            ["是否存在人数、城市、天数或特殊业务场景差异？"],
                            ["补充数量/人数/期间等标准化口径并核对原始单据"],
                            [{"factor": "robust_outlier", "points": 2, "robust_z": round(robust_z, 4), "peer_group": list(key)}])
    if valid_groups == 0:
        skipped.append("robust-outlier：没有达到最小样本且 MAD 非零的 peer group")
    return skipped, {
        "near_duplicate_window_days": near_days,
        "near_duplicate_amount_tolerance": near_tolerance,
        "split_window_days": split_days,
        "weekend_check": bool(policy.get("weekend_check", False)),
        "outlier_min_group_size": min_group,
        "outlier_robust_z": z_threshold,
        "policy_version": policy.get("policy_version"),
    }


def write_csv(path: Path, rows: Sequence[Dict[str, Any]], fields: Sequence[str]) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fields), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_jsonl(path: Path, rows: Sequence[Dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def validate(findings: Sequence[Dict[str, Any]], evidence: Sequence[Dict[str, Any]]) -> None:
    evidence_ids = {item["evidence_id"] for item in evidence}
    finding_ids = set()
    for finding in findings:
        if finding["finding_id"] in finding_ids:
            raise ValueError("finding_id 重复")
        finding_ids.add(finding["finding_id"])
        if not finding["evidence_refs"] or not set(finding["evidence_refs"]).issubset(evidence_ids):
            raise ValueError("finding %s 存在缺失 evidence 引用" % finding["finding_id"])
        if not finding.get("human_review_required"):
            raise ValueError("finding 必须要求人工复核")


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--policy", type=Path)
    parser.add_argument("--field-map", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--check-env", action="store_true")
    args = parser.parse_args(argv)
    if args.check_env:
        try:
            import openpyxl
            xlsx = "available (%s)" % openpyxl.__version__
        except ImportError:
            xlsx = "unavailable; CSV mode still works"
        print(json.dumps({"python": sys.version.split()[0], "recommended": ">=3.10", "xlsx_support": xlsx, "network_required": False}, ensure_ascii=False))
        return 0
    if not args.input or not args.output:
        parser.error("正式运行需要 --input 和 --output")
    input_path = args.input.resolve()
    if not input_path.is_file():
        raise FileNotFoundError(input_path)
    output = args.output.resolve()
    if output.exists() and any(output.iterdir()):
        raise FileExistsError("输出目录非空；请指定新的目录：%s" % output)
    output.mkdir(parents=True, exist_ok=True)
    started = datetime.now(timezone.utc)
    source_hash = sha256_file(input_path)
    policy = load_json(args.policy.resolve() if args.policy else None)
    explicit_map = load_json(args.field_map.resolve() if args.field_map else None)
    headers, source_rows, sheet = read_table(input_path, args.sheet)
    mapping, warnings = resolve_mapping(headers, explicit_map)
    if not policy.get("default_currency"):
        warnings.append("未配置 default_currency，缺失币种按 CNY 处理")
    clean, bad, changes = normalize_rows(source_rows, mapping, str(policy.get("default_currency", "CNY")))
    builder = ResultBuilder(input_path.name, source_hash)
    skipped, parameters = run_rules(clean, policy, builder)
    validate(builder.findings, builder.evidence)

    write_csv(output / "clean_expenses.csv", clean, OUTPUT_FIELDS + ("_source_row", "_source_sheet"))
    write_csv(output / "bad_rows.csv", bad, ("source_row", "reasons", "raw_record"))
    flat = []
    for finding in builder.findings:
        flat.append({
            "finding_id": finding["finding_id"], "finding_type": finding["finding_type"],
            "title": finding["title"], "risk_priority": finding["risk_priority"],
            "evidence_strength": finding["evidence_strength"], "risk_score": finding["risk_score"],
            "entities": json.dumps(finding["entities"], ensure_ascii=False),
            "evidence_refs": "|".join(finding["evidence_refs"]),
        })
    write_csv(output / "findings.csv", flat, ("finding_id", "finding_type", "title", "risk_priority", "evidence_strength", "risk_score", "entities", "evidence_refs"))
    write_jsonl(output / "findings.jsonl", builder.findings)
    write_jsonl(output / "evidence.jsonl", builder.evidence)

    null_rates = {}
    for header in headers:
        null_rates[header] = round(sum(1 for row in source_rows if not norm_text(row.get(header))) / max(len(source_rows), 1), 6)
    quality = [
        "# 数据质量报告", "", "- 源文件：`%s`" % input_path.name,
        "- 工作表：`%s`" % (sheet or "CSV"), "- 源数据行数：%d" % len(source_rows),
        "- 有效行数：%d" % len(clean), "- Bad rows：%d" % len(bad),
        "- 字段映射：`%s`" % json.dumps(mapping, ensure_ascii=False),
        "- 标准化计数：`%s`" % json.dumps(changes, ensure_ascii=False), "", "## 源字段空值率", "",
    ]
    quality.extend("- `%s`: %.2f%%" % (field, rate * 100) for field, rate in null_rates.items())
    if warnings or skipped:
        quality.extend(["", "## 警告与跳过规则", ""] + ["- " + item for item in warnings + skipped])
    (output / "data_quality.md").write_text("\n".join(quality) + "\n", encoding="utf-8")

    counts = defaultdict(int)
    priorities = defaultdict(int)
    for finding in builder.findings:
        counts[finding["finding_type"]] += 1
        priorities[finding["risk_priority"]] += 1
    summary = [
        "# 费用审计确定性摘要", "", "- 分析有效记录：%d；排除坏行：%d。" % (len(clean), len(bad)),
        "- Findings：%d；Evidence：%d。" % (len(builder.findings), len(builder.evidence)),
        "- 风险优先级：`%s`" % json.dumps(dict(priorities), ensure_ascii=False, sort_keys=True),
        "- 发现类型：`%s`" % json.dumps(dict(counts), ensure_ascii=False, sort_keys=True), "",
        "> 这些结果是复核线索，不是舞弊、虚假报销或拒付结论。请先查看数据质量，再回到 evidence 对应源行核验。", "",
        "## 建议复核顺序", "",
        "1. 先查 high/critical 且 evidence_strength 为 strong 的重复或制度命中。",
        "2. 再按员工、商户和期间合并 near duplicate、split 和 outlier 模式。",
        "3. 周末信号单独保持低优先级，主动寻找值班、出差和客户现场等合理解释。",
    ]
    (output / "summary.md").write_text("\n".join(summary) + "\n", encoding="utf-8")

    finished = datetime.now(timezone.utc)
    script_path = Path(__file__).resolve()
    input_files = [{"path": input_path.name, "sha256": source_hash, "sheet": sheet}]
    if args.policy:
        input_files.append({"path": args.policy.name, "sha256": sha256_file(args.policy.resolve())})
    if args.field_map:
        input_files.append({"path": args.field_map.name, "sha256": sha256_file(args.field_map.resolve())})
    manifest = {
        "run_id": "EXP-%s-%s" % (started.strftime("%Y%m%dT%H%M%SZ"), source_hash[:8]),
        "skill": SKILL, "skill_version": VERSION, "started_at": started.isoformat(), "finished_at": finished.isoformat(),
        "input_files": input_files, "parameters": parameters, "field_mapping": mapping,
        "scripts": {script_path.name: "sha256:" + sha256_file(script_path)},
        "warnings": warnings, "skipped_rules": skipped, "network_access": False,
        "outputs": ["clean_expenses.csv", "bad_rows.csv", "findings.csv", "findings.jsonl", "evidence.jsonl", "summary.md", "data_quality.md", "run_manifest.json"],
    }
    (output / "run_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(output), "valid_rows": len(clean), "bad_rows": len(bad), "findings": len(builder.findings), "evidence": len(builder.evidence)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print("ERROR: %s" % exc, file=sys.stderr)
        raise SystemExit(2)
